import os
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, PatternFill, Border, Side, Font
from openpyxl.drawing.image import Image as XLImage
from webapp.config import Config

def save_to_excel(rows, path_excel, max_secondary=10, max_layouts=10):
    # Salva i dati in un file Excel con formattazione avanzata:
    # - Pulizia dei valori (N/A)
    # - Inserimento delle immagini ridimensionate
    # - Hyperlink cliccabili
    # - Colonne dinamiche per foto e layout
    # - Evidenziazione delle celle mancanti

    # Converte i dati in DataFrame
    df = pd.DataFrame(rows)

    # Funzione per pulire valori vuoti o invalidi
    def clean_and_fill(val):
        # Normalizza valori vuoti, None, NaN, stringhe "none"
        if pd.isna(val) or val is None or str(val).strip() == "" or str(val).lower() == "none":
            return "N/A"
        # Rimuove eventuali doppi apici e spazi
        return str(val).replace('"', '').strip()

    # Applica la pulizia a tutte le colonne tranne il codice
    for col in df.columns:
        if col != "code":
            df[col] = df[col].apply(clean_and_fill)

    # Costruzione dinamica delle colonne richieste
    # PDF principali
    pdf_links = ["manual_pdf", "safety_pdf"]

    # Colonne del layout dinamiche (layout_1 ... layout_N)
    layout_cols = [f"layout_{i}" for i in range(1, max_layouts + 1)]

    # Colonne delle foto secondarie (photo_1 ... photo_N)
    photo_cols = [f"photo_{i}" for i in range(1, max_secondary + 1)]

    # Colonne che contengono le immagini (principale + secondarie + layout)
    image_headers = ["image_main"] + photo_cols + layout_cols

    # Assicura che tutte le colonne esistano nel DataFrame
    for col in (pdf_links + layout_cols + photo_cols):
        if col not in df.columns:
            df[col] = "N/A"

    # Ordine finale delle colonne nel file Excel
    columns_order = [
        "code", "name", "scale", "image_main",
        *photo_cols,
        *layout_cols,
        *pdf_links,
        "dimensions", "epoch", "kit_contains",
        "lighting", "category", "ean", "product_link"
    ]

    # Mantiene solo le colonne effettivamente presenti
    df = df[[c for c in columns_order if c in df.columns]]

    # Esporta il DataFrame in Excel (prima fase, senza formattazione)
    df.to_excel(path_excel, index=False)

    # FORMATTZIONE AVANZATA OPENPYXL
    wb = load_workbook(path_excel)
    ws = wb.active

    # Stile delle celle
    red_fill = PatternFill(start_color="FFCCCC", end_color="FFCCCC", fill_type="solid")
    center_style = Alignment(horizontal="center", vertical="center", wrap_text=True)
    border_style = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin')
    )

    # Dimensioni massime delle immagini (ridimensionamento proporzionale)
    MAX_W, MAX_H = 180, 170

    # Iterazione su tutte le celle per applicare la formattazione e le immagini
    for row_idx, row in enumerate(ws.iter_rows(min_row=1), 1):

        # Altezza delle righe: header più basso, righe dei dati più alte
        ws.row_dimensions[row_idx].height = 140 if row_idx > 1 else 30

        for cell in row:
            header = ws.cell(row=1, column=cell.column).value
            cell_val = str(cell.value).strip()

            # Applica lo stile base
            cell.alignment = center_style
            cell.border = border_style

            # Evidenzia i valori mancanti
            if row_idx > 1 and cell_val == "N/A":
                cell.fill = red_fill

            # INSERIMENTO DELLE IMMAGINI (image_main, photo_X, layout_X)
            if row_idx > 1 and header in image_headers:

                if cell_val != "N/A":
                    # Costruisce il percorso assoluto dell'immagine
                    img_path = os.path.join(Config.IMG_DIR, cell_val)

                    if os.path.exists(img_path):
                        cell.value = "" # Rimuove il testo per lasciare spazio all'immagine

                        try:
                            img = XLImage(img_path)

                            # Ridimensionamento proporzionale
                            ratio = min(MAX_W / img.width, MAX_H / img.height)
                            img.width = int(img.width * ratio)
                            img.height = int(img.height * ratio)

                            # Inserisce l'immagine nella cella
                            ws.add_image(img, cell.coordinate)

                        except:
                            # In caso di errore nell'immagine
                            cell.value = "Errore Img"
                    else:
                        # File immagine non trovato
                        cell.value = "N/A"
                        cell.fill = red_fill

            # HYPERLINK CLICCABILI (product_link, manual_pdf, safety_pdf)
            elif row_idx > 1 and "http" in cell_val and header in ["product_link", "manual_pdf", "safety_pdf"]:

                # Etichetta diversa per il link al sito e per il PDF
                label = "VAI AL SITO" if header == "product_link" else "APRI PDF"

                # Formula di Excel per l'hyperlink
                cell.value = f'=HYPERLINK("{cell_val}", "{label}")'
                cell.font = Font(color="0000FF", underline="single")

    # Larghezza colonne dinamica:
    # - Colonne immagini più larghe
    # - Colonne testo standard più strette
    for col in ws.columns:
        header = ws.cell(row=1, column=col[0].column).value

        if header == "name":
            ws.column_dimensions[col[0].column_letter].width = 60
        elif header in image_headers:
            ws.column_dimensions[col[0].column_letter].width = 28
        else:
            ws.column_dimensions[col[0].column_letter].width = 20

    # Salva il file Excel finale
    wb.save(path_excel)